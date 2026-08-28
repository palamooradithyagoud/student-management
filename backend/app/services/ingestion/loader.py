import re
from pathlib import Path
from typing import Optional, List, Dict, Any
import pandas as pd
import openpyxl
import xlrd
from backend.app.services.ingestion.inspector import FileInspector, match_column_to_canonical

class RawDataLoader:
    """
    Loads raw Excel/CSV files and converts them into standardized DataFrames
    handling:
    1. Multi-level matrix Result sheets with (Code) Name and GD/S/GP sub-columns.
    2. Multi-level matrix Attendance sheets with Conducted (C) and Attended (A) sub-columns.
    3. Generic tabular CSV/Excel files.
    """

    @staticmethod
    def load_dataset(file_path: Path, dataset_type: str, semester: int) -> pd.DataFrame:
        ext = file_path.suffix.lower()
        
        # Check if this is an authentic multi-level Result file
        if "result" in dataset_type.lower() or "result" in file_path.name.lower():
            try:
                df_res = RawDataLoader.load_matrix_results(file_path, semester)
                if not df_res.empty:
                    df_res["dataset_type"] = dataset_type
                    return df_res
            except Exception as e:
                # Fall back to standard parser if special parser encounters a different schema
                pass

        # Check if this is an authentic multi-level Attendance file
        if "attendance" in dataset_type.lower() or "attendance" in file_path.name.lower():
            try:
                df_att = RawDataLoader.load_matrix_attendance(file_path, semester)
                if not df_att.empty:
                    df_att["dataset_type"] = dataset_type
                    return df_att
            except Exception as e:
                # Fall back to standard parser
                pass

        # Standard loader fallback
        inspection = FileInspector.inspect_file(file_path)
        header_idx = inspection["detected_header_row"]

        if ext in [".xlsx", ".xls"]:
            df_raw = pd.read_excel(file_path, sheet_name=0, skiprows=header_idx)
        else:
            try:
                df_raw = pd.read_csv(file_path, skiprows=header_idx, encoding='utf-8')
            except UnicodeDecodeError:
                df_raw = pd.read_csv(file_path, skiprows=header_idx, encoding='latin1')

        df_raw = df_raw.dropna(how='all')

        col_map = {}
        for col in df_raw.columns:
            canon = match_column_to_canonical(str(col))
            if canon:
                col_map[col] = canon

        df = df_raw.rename(columns=col_map)

        has_roll = "roll_no" in df.columns
        has_subject = "subject_code" in df.columns or "subject_name" in df.columns

        if has_roll and not has_subject:
            df = RawDataLoader._melt_wide_format(df, dataset_type, semester)
        else:
            df["semester"] = semester
            df["dataset_type"] = dataset_type

        return df

    @staticmethod
    def load_matrix_results(file_path: Path, semester: int) -> pd.DataFrame:
        """
        Parses authentic college result sheets with:
        - Row 1: Subject Codes & Names in merged/spaced columns e.g. '( A9001 ) MAC'
        - Row 2: Sub-headers 'GD' (Grade), 'S' (Status), 'GP' (Grade Points), 'SGPA', 'No. of Backlog'
        - Row 3+: Student rows with S.No, Student Name, Roll No, Section, and per-subject marks.
        """
        ext = file_path.suffix.lower()
        col_subjects: Dict[int, tuple] = {}
        roll_col_idx = 2
        name_col_idx = 1
        sec_col_idx = 3
        sgpa_col_idx = None
        backlog_col_idx = None

        rows_data = []

        if ext == ".xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active

            current_sub_code = None
            current_sub_name = None

            for col in range(1, ws.max_column + 1):
                v1 = ws.cell(row=1, column=col).value
                v2 = ws.cell(row=2, column=col).value

                v1_str = str(v1).strip() if v1 is not None else ""
                v2_str = str(v2).strip() if v2 is not None else ""
                combined_hdr = f"{v1_str} {v2_str}".strip().lower()

                if "roll" in combined_hdr:
                    roll_col_idx = col
                elif "student name" in combined_hdr or "name" in v2_str.lower():
                    name_col_idx = col
                elif "sec" in combined_hdr:
                    sec_col_idx = col
                elif "sgpa" in combined_hdr:
                    sgpa_col_idx = col
                elif "backlog" in combined_hdr:
                    backlog_col_idx = col

                if v1_str:
                    if "FINAL" in v1_str.upper() or "STUDENT" in v1_str.upper() or "INFO" in v1_str.upper():
                        current_sub_code = None
                        current_sub_name = None
                    else:
                        m = re.search(r'\(\s*([A-Z0-9]+)\s*\)\s*(.*)', v1_str, re.IGNORECASE)
                        if m:
                            current_sub_code = m.group(1).strip().upper()
                            current_sub_name = m.group(2).strip()
                        else:
                            current_sub_code = v1_str.upper()
                            current_sub_name = v1_str

                if current_sub_code and v2_str:
                    col_subjects[col] = (current_sub_code, current_sub_name, v2_str.upper())

            for r in range(3, ws.max_row + 1):
                roll = ws.cell(row=r, column=roll_col_idx).value
                if not roll or str(roll).strip() == "":
                    continue

                roll_clean = str(roll).strip().upper()
                name_clean = str(ws.cell(row=r, column=name_col_idx).value or "").strip()
                sec_clean = str(ws.cell(row=r, column=sec_col_idx).value or "A").strip()
                sgpa_val = ws.cell(row=r, column=sgpa_col_idx).value if sgpa_col_idx else None
                
                raw_backlog = ws.cell(row=r, column=backlog_col_idx).value if backlog_col_idx else None
                try:
                    backlog_val = int(float(str(raw_backlog).strip())) if (raw_backlog is not None and str(raw_backlog).strip() != "") else 0
                except (ValueError, TypeError):
                    backlog_val = 0

                # Per-subject readings
                sub_dict: Dict[str, Dict[str, Any]] = {}
                for col, (sc, sn, stype) in col_subjects.items():
                    if sc not in sub_dict:
                        sub_dict[sc] = {"code": sc, "name": sn, "GD": None, "S": None, "GP": None}
                    sub_dict[sc][stype] = ws.cell(row=r, column=col).value

                for sc, svals in sub_dict.items():
                    status_raw = str(svals.get("S") or "").strip().upper()
                    status_clean = "PASS" if status_raw == "P" else ("FAIL" if status_raw == "F" else (status_raw or "PASS"))
                    
                    gp_raw = svals.get("GP")
                    gp_clean = float(gp_raw) if (gp_raw is not None and str(gp_raw).strip() != "") else None

                    # If grade points available, estimate approximate marks out of 100
                    marks_clean = float(gp_clean * 10.0) if gp_clean is not None else None

                    rows_data.append({
                        "roll_no": roll_clean,
                        "student_name": name_clean if name_clean else None,
                        "section": sec_clean,
                        "semester": semester,
                        "subject_code": sc,
                        "subject_name": svals.get("name") or sc,
                        "grade": str(svals.get("GD") or "").strip() or None,
                        "grade_points": gp_clean,
                        "marks": marks_clean,
                        "status": status_clean,
                        "sgpa": float(sgpa_val) if (sgpa_val is not None and str(sgpa_val).strip() != "") else None,
                        "backlog_count": backlog_val
                    })

        elif ext == ".xls":
            wb = xlrd.open_workbook(file_path)
            sh = wb.sheet_by_index(0)

            current_sub_code = None
            current_sub_name = None

            for col in range(sh.ncols):
                v1 = sh.cell_value(0, col)
                v2 = sh.cell_value(1, col)

                v1_str = str(v1).strip() if v1 != "" else ""
                v2_str = str(v2).strip() if v2 != "" else ""
                combined_hdr = f"{v1_str} {v2_str}".strip().lower()

                if "roll" in combined_hdr:
                    roll_col_idx = col
                elif "student name" in combined_hdr or "name" in v2_str.lower():
                    name_col_idx = col
                elif "sec" in combined_hdr:
                    sec_col_idx = col
                elif "sgpa" in combined_hdr:
                    sgpa_col_idx = col
                elif "backlog" in combined_hdr:
                    backlog_col_idx = col

                if v1_str:
                    if "FINAL" in v1_str.upper() or "STUDENT" in v1_str.upper() or "INFO" in v1_str.upper():
                        current_sub_code = None
                        current_sub_name = None
                    else:
                        m = re.search(r'\(\s*([A-Z0-9]+)\s*\)\s*(.*)', v1_str, re.IGNORECASE)
                        if m:
                            current_sub_code = m.group(1).strip().upper()
                            current_sub_name = m.group(2).strip()
                        else:
                            current_sub_code = v1_str.upper()
                            current_sub_name = v1_str

                if current_sub_code and v2_str:
                    col_subjects[col] = (current_sub_code, current_sub_name, v2_str.upper())

            for r in range(2, sh.nrows):
                roll = sh.cell_value(r, roll_col_idx)
                if not roll or str(roll).strip() == "":
                    continue

                roll_clean = str(roll).strip().upper()
                name_clean = str(sh.cell_value(r, name_col_idx) or "").strip()
                sec_clean = str(sh.cell_value(r, sec_col_idx) or "A").strip()
                sgpa_val = sh.cell_value(r, sgpa_col_idx) if sgpa_col_idx is not None else None
                
                raw_backlog = sh.cell_value(r, backlog_col_idx) if backlog_col_idx is not None else None
                try:
                    backlog_val = int(float(str(raw_backlog).strip())) if (raw_backlog is not None and str(raw_backlog).strip() != "") else 0
                except (ValueError, TypeError):
                    backlog_val = 0

                sub_dict = {}
                for col, (sc, sn, stype) in col_subjects.items():
                    if sc not in sub_dict:
                        sub_dict[sc] = {"code": sc, "name": sn, "GD": None, "S": None, "GP": None}
                    sub_dict[sc][stype] = sh.cell_value(r, col)

                for sc, svals in sub_dict.items():
                    status_raw = str(svals.get("S") or "").strip().upper()
                    status_clean = "PASS" if status_raw == "P" else ("FAIL" if status_raw == "F" else (status_raw or "PASS"))
                    
                    gp_raw = svals.get("GP")
                    gp_clean = float(gp_raw) if (gp_raw is not None and str(gp_raw).strip() != "") else None
                    marks_clean = float(gp_clean * 10.0) if gp_clean is not None else None

                    rows_data.append({
                        "roll_no": roll_clean,
                        "student_name": name_clean if name_clean else None,
                        "section": sec_clean,
                        "semester": semester,
                        "subject_code": sc,
                        "subject_name": svals.get("name") or sc,
                        "grade": str(svals.get("GD") or "").strip() or None,
                        "grade_points": gp_clean,
                        "marks": marks_clean,
                        "status": status_clean,
                        "sgpa": float(sgpa_val) if (sgpa_val is not None and str(sgpa_val).strip() != "") else None,
                        "backlog_count": backlog_val
                    })

        return pd.DataFrame(rows_data)

    @staticmethod
    def load_matrix_attendance(file_path: Path, semester: int) -> pd.DataFrame:
        """
        Parses authentic college attendance sheets with:
        - Row 4: Subject Headers e.g. 'A9001 (MAC)' or 'A9002\\nODECV'
        - Row 6: Conducted (C) and Attended (A) indicators
        - Row 7+: Student records with Roll Number and subject attendance counts.
        """
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        sub_col_map: Dict[int, tuple] = {}
        cur_sub_code = None
        cur_sub_name = None

        roll_col = 2
        name_col = None
        sec_col = None

        for c in range(1, ws.max_column + 1):
            v4 = ws.cell(row=4, column=c).value
            if v4:
                v4_str = str(v4).strip()
                if "Roll" in v4_str:
                    roll_col = c
                elif "Name" in v4_str:
                    name_col = c
                elif "Section" in v4_str or "Sec" in v4_str:
                    sec_col = c
                elif "Total" in v4_str:
                    cur_sub_code = None
                else:
                    parts = v4_str.replace("\n", " ").split("(")
                    cur_sub_code = parts[0].strip().upper()
                    cur_sub_name = parts[1].replace(")", "").strip() if len(parts) > 1 else cur_sub_code

            c_type = ws.cell(row=6, column=c).value
            if cur_sub_code and c_type:
                sub_col_map[c] = (cur_sub_code, cur_sub_name, str(c_type).strip().upper())

        sec_hint = "A" if ("_A" in file_path.name or "- A" in file_path.name) else ("B" if ("_B" in file_path.name or "- B" in file_path.name) else "C")

        records = []
        for r in range(7, ws.max_row + 1):
            roll = ws.cell(row=r, column=roll_col).value
            if not roll or str(roll).strip() == "":
                continue

            roll_clean = str(roll).strip().upper()
            name_clean = str(ws.cell(row=r, column=name_col).value or "").strip() if name_col else None
            sec_clean = str(ws.cell(row=r, column=sec_col).value or sec_hint).strip() if sec_col else sec_hint

            sub_att: Dict[str, Dict[str, Any]] = {}
            for c, (sc, sn, ctype) in sub_col_map.items():
                if sc not in sub_att:
                    sub_att[sc] = {"code": sc, "name": sn, "C": 0, "A": 0}
                val = ws.cell(row=r, column=c).value
                if val is not None and str(val).isdigit():
                    sub_att[sc][ctype] = float(val)

            for sc, avals in sub_att.items():
                cond = avals["C"]
                att = avals["A"]
                pct = round((att / cond * 100.0), 2) if cond > 0 else 0.0
                records.append({
                    "roll_no": roll_clean,
                    "student_name": name_clean if name_clean else None,
                    "section": sec_clean,
                    "semester": semester,
                    "subject_code": sc,
                    "subject_name": avals.get("name") or sc,
                    "attendance_percentage": pct,
                    "classes_attended": int(att),
                    "classes_conducted": int(cond)
                })

        return pd.DataFrame(records)

    @staticmethod
    def _melt_wide_format(df: pd.DataFrame, dataset_type: str, semester: int) -> pd.DataFrame:
        id_vars = [c for c in ["roll_no", "student_name", "section", "batch", "sgpa", "backlog_count"] if c in df.columns]
        value_vars = [c for c in df.columns if c not in id_vars and str(c).strip() != ""]

        val_col_name = "attendance_percentage" if "attendance" in dataset_type else "marks"

        melted = pd.melt(
            df,
            id_vars=id_vars,
            value_vars=value_vars,
            var_name="raw_subject_indicator",
            value_name=val_col_name
        )

        melted["semester"] = semester
        melted["dataset_type"] = dataset_type
        melted["subject_code"] = melted["raw_subject_indicator"].astype(str).str.strip().str.upper()
        melted["subject_name"] = melted["raw_subject_indicator"].astype(str).str.strip()

        return melted
