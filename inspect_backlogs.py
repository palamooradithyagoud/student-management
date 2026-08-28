import openpyxl
import xlrd
import pandas as pd

wb1 = openpyxl.load_workbook('data/sem result/RESULT 1-1.xlsx', data_only=True)
sheet1 = wb1.active
print('=== RESULT 1-1 Columns ===')
for c in range(1, sheet1.max_column + 1):
    val1 = sheet1.cell(1, c).value
    val2 = sheet1.cell(2, c).value
    print(f'Col {c}: Row1="{val1}", Row2="{val2}"')

print('\n=== Sample Row 3 to 8 of RESULT 1-1 (end columns) ===')
for r in range(3, 9):
    row_vals = [f'{sheet1.cell(2, c).value or sheet1.cell(1, c).value}: {sheet1.cell(r, c).value}' for c in range(sheet1.max_column - 10, sheet1.max_column + 1)]
    print(f'Row {r} ({sheet1.cell(r, 3).value}):', ', '.join(row_vals))

print('\n=== Results 1-2 Columns ===')
wb2 = xlrd.open_workbook('data/sem result/Results 1-2.xls')
sheet2 = wb2.sheet_by_index(0)
for c in range(sheet2.ncols):
    val0 = sheet2.cell_value(0, c)
    val1 = sheet2.cell_value(1, c)
    val2 = sheet2.cell_value(2, c)
    print(f'Col {c+1}: Row0="{val0}", Row1="{val1}", Row2="{val2}"')

print('\n=== Sample Row 3 to 8 of Results 1-2 (end columns) ===')
for r in range(3, 9):
    row_vals = [f'{sheet2.cell_value(1, c) or sheet2.cell_value(0, c)}: {sheet2.cell_value(r, c)}' for c in range(sheet2.ncols - 10, sheet2.ncols)]
    print(f'Row {r} ({sheet2.cell_value(r, 2)}):', ', '.join(row_vals))
